# backend/app.py
from flask import Flask, request, jsonify
from flask_cors import CORS
import os, io
from model_loader import predict_fake_news
from PyPDF2 import PdfReader
from docx import Document
from PIL import Image
import pytesseract
import openpyxl

app = Flask(__name__)
CORS(app)

@app.route('/')
def home():
    return jsonify({"message": "Fake News Detector API is running ✅"})

# --- Analyze Text Endpoint ---
@app.route('/analyze-text', methods=['POST'])
def analyze_text():
    data = request.get_json()
    if not data or 'text' not in data:
        return jsonify({"error": "Missing text field"}), 400

    text = data['text'].strip()
    if not text:
        return jsonify({"error": "Empty text"}), 400

    result = predict_fake_news(text)
    result['preview'] = text[:300]
    return jsonify(result)

# --- Analyze File Endpoint ---
@app.route('/analyze-file', methods=['POST'])
def analyze_file():
    if 'file' not in request.files:
        return jsonify({"error": "No file provided"}), 400

    file = request.files['file']
    filename = file.filename.lower()
    text = ""

    try:
        if filename.endswith('.pdf'):
            reader = PdfReader(file)
            for page in reader.pages:
                text += page.extract_text() or ""
        elif filename.endswith(('.doc', '.docx')):
            doc = Document(file)
            for para in doc.paragraphs:
                text += para.text + "\n"
        elif filename.endswith(('.jpg', '.jpeg', '.png')):
            image = Image.open(file)
            text = pytesseract.image_to_string(image)
        elif filename.endswith('.xlsx'):
            wb = openpyxl.load_workbook(file)
            for sheet in wb.worksheets:
                for row in sheet.iter_rows(values_only=True):
                    text += " ".join([str(cell) for cell in row if cell]) + "\n"
        else:
            return jsonify({"error": "Unsupported file type"}), 400

        if not text.strip():
            return jsonify({"error": "No readable text found in file"}), 400

        result = predict_fake_news(text)
        result['preview'] = text[:300]
        return jsonify(result)

    except Exception as e:
        print("Error processing file:", e)
        return jsonify({"error": f"Failed to process file: {str(e)}"}), 500


if __name__ == '__main__':
    port = int(os.environ.get("PORT", 5001))
    app.run(host='0.0.0.0', port=port, debug=True)
